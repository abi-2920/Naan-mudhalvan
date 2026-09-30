import streamlit as st
import google.generativeai as genai
from docx import Document
import io

# 1. Page Configuration
st.set_page_config(page_title="AI Legal Document Generator", page_icon="⚖️", layout="wide")

st.title("⚖️ AI-Powered Legal Document Generator")
st.write("Generate customized legal agreements instantly using AI.")

# 2. Sidebar API Key Setup
st.sidebar.header("Configuration")
api_key = st.sidebar.text_input("Enter Gemini API Key", type="password")

if api_key:
    genai.configure(api_key=api_key)

# 3. Inputs
st.subheader("1. Select Document Type")
doc_type = st.selectbox("Document Type", [
    "Rental Agreement", 
    "Non-Disclosure Agreement (NDA)", 
    "Employment Contract", 
    "Freelance Service Agreement"
])

st.subheader("2. Enter Agreement Details")
col1, col2 = st.columns(2)

with col1:
    party1_name = st.text_input("Party 1 Name (e.g., Owner/Employer)")
    party1_address = st.text_area("Party 1 Address")

with col2:
    party2_name = st.text_input("Party 2 Name (e.g., Tenant/Employee)")
    party2_address = st.text_area("Party 2 Address")

specific_terms = st.text_area("Specific Terms & Conditions (e.g., Rent Amount, Duration, Notice Period)")

# Function to generate DOCX
def create_docx(text):
    doc = Document()
    doc.add_heading(f'{doc_type}', level=1)
    doc.add_paragraph(text)
    
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# 4. Generate Action
if st.button("Generate Legal Document"):
    if not api_key:
        st.error("Please enter your Gemini API Key in the sidebar.")
    elif not party1_name or not party2_name:
        st.warning("Please fill in both Party names.")
    else:
        with st.spinner("Generating legal document with AI..."):
            try:
                model = genai.GenerativeModel('gemini-3.6-flash') 
                prompt = f"""
                Act as an expert legal consultant. Draft a professional, legally binding {doc_type} using Indian Legal Standards.
                
                Details:
                - Party 1: {party1_name}, Address: {party1_address}
                - Party 2: {party2_name}, Address: {party2_address}
                - Specific Terms & Rules: {specific_terms}
                
                Format the output cleanly with proper Headings, Sections, Clauses, Signature Blocks, and Disclaimer.
                """
                
                response = model.generate_content(prompt)
                generated_text = response.text
                
                st.success("Document Generated Successfully!")
                st.text_area("Preview Generated Document", generated_text, height=400)
                
                # File Download Option
                docx_data = create_docx(generated_text)
                st.download_button(
                    label="📥 Download as Word (.docx)",
                    data=docx_data,
                    file_name=f"{doc_type.replace(' ', '_')}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
            except Exception as e:
                st.error(f"Error generating document: {e}")

